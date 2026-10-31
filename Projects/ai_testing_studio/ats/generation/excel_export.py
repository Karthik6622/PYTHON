"""Excel export (openpyxl). Cells that could be read as formulas are neutralised (CSV/Excel injection)."""

from __future__ import annotations

import io
from collections import Counter
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ats.models import Scenario, TestCase, utc_now
from ats.repository import Project
from ats.security import mask

HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SCENARIO_COLUMNS = [
    ("Scenario ID", "scenario_id", 16), ("Module", "module", 18), ("Page", "page", 36), ("Title", "title", 40),
    ("Description", "description", 50), ("Preconditions", "preconditions", 36), ("Test Type", "test_type", 14),
    ("Priority", "priority", 10), ("Expected Outcome", "expected_outcome", 46), ("Automation Suitability", "automation_suitability", 16),
    ("Status", "status", 11),
]  # fmt: skip
TEST_CASE_COLUMNS = [
    ("Test Case ID", "test_case_id", 16), ("Scenario ID", "scenario_id", 16), ("Module", "module", 16), ("Title", "title", 40),
    ("Description", "description", 44), ("Preconditions", "preconditions", 32), ("Test Data", "test_data", 32), ("Steps", "steps", 60),
    ("Expected Result", "expected_result", 44), ("Test Type", "test_type", 13), ("Priority", "priority", 10),
    ("Automation Candidate", "automation_candidate", 12), ("Related Page", "related_page", 34), ("Component", "component", 18),
    ("Tags", "tags", 20), ("Status", "status", 11),
]  # fmt: skip


def safe_cell(value: Any) -> Any:
    """Neutralise formula injection and secrets; lists and dicts become readable text."""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, list):
        value = "\n".join(f"{i}. {v}" for i, v in enumerate(value, 1)) if len(value) > 1 else "".join(map(str, value))
    elif isinstance(value, dict):
        value = "\n".join(f"{k}: {v}" for k, v in value.items())
    text = mask(str(value if value is not None else ""))
    return "'" + text if text[:1] in {"=", "+", "-", "@", "\t", "\r"} else text[:32000]


def _write_table(ws: Worksheet, columns: list[tuple[str, str, int]], rows: list[dict[str, Any]]) -> None:
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([safe_cell(row.get(key)) for _, key, _ in columns])
    for index, (_, _, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _scenario_sheet(wb: Workbook, scenarios: list[Scenario], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = "Test Scenarios"
    _write_table(ws, SCENARIO_COLUMNS, [s.model_dump() for s in scenarios])


def _case_sheet(wb: Workbook, cases: list[TestCase], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = "Test Cases"
    _write_table(ws, TEST_CASE_COLUMNS, [c.model_dump() for c in cases])


def coverage_summary(project: Project, scenarios: list[Scenario], cases: list[TestCase]) -> list[tuple[str, Any]]:
    """Flat label/value rows used by the Excel sheet and the UI."""
    by_scenario = Counter(c.scenario_id for c in cases)
    uncovered = [s.scenario_id for s in scenarios if by_scenario.get(s.scenario_id, 0) == 0 and s.status == "approved"]
    rows: list[tuple[str, Any]] = [
        ("Project", project.name), ("Application URL", project.url), ("Environment", project.environment),
        ("Coverage level", project.coverage), ("Generated at", utc_now()),
        ("Total scenarios", len(scenarios)), ("Approved scenarios", sum(s.status == "approved" for s in scenarios)),
        ("Total test cases", len(cases)), ("Approved test cases", sum(c.status == "approved" for c in cases)),
        ("Automation candidates", sum(c.automation_candidate for c in cases)),
        ("Approved scenarios without test cases", ", ".join(uncovered) or "None"),
    ]  # fmt: skip
    for title, counter in (
        ("Scenarios by module", Counter(s.module for s in scenarios)),
        ("Scenarios by test type", Counter(s.test_type for s in scenarios)),
        ("Scenarios by priority", Counter(s.priority for s in scenarios)),
        ("Test cases by module", Counter(c.module for c in cases)),
        ("Test cases by type", Counter(c.test_type for c in cases)),
        ("Test cases by priority", Counter(c.priority for c in cases)),
    ):
        rows.append((title, ""))
        rows.extend((f"    {name}", count) for name, count in sorted(counter.items()))
    return rows


def _summary_sheet(wb: Workbook, project: Project, scenarios: list[Scenario], cases: list[TestCase]) -> None:
    ws = wb.create_sheet("Coverage Summary")
    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for label, value in coverage_summary(project, scenarios, cases):
        ws.append([safe_cell(label) if not label.startswith("    ") else label, safe_cell(value)])
        if value == "" and not label.startswith("    "):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 60


def _bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def scenarios_workbook(project: Project, scenarios: list[Scenario], cases: list[TestCase] | None = None) -> bytes:
    wb = Workbook()
    _scenario_sheet(wb, scenarios, first=True)
    _summary_sheet(wb, project, scenarios, cases or [])
    return _bytes(wb)


def cases_workbook(project: Project, scenarios: list[Scenario], cases: list[TestCase]) -> bytes:
    wb = Workbook()
    _case_sheet(wb, cases, first=True)
    _summary_sheet(wb, project, scenarios, cases)
    return _bytes(wb)


def combined_workbook(project: Project, scenarios: list[Scenario], cases: list[TestCase]) -> bytes:
    """Test Scenarios, Test Cases and Coverage Summary in a single workbook."""
    wb = Workbook()
    _scenario_sheet(wb, scenarios, first=True)
    _case_sheet(wb, cases)
    _summary_sheet(wb, project, scenarios, cases)
    return _bytes(wb)
