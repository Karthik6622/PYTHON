"""Build the book DOCX with python-docx (6x9 inch pages, heading styles, table of contents field)."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage

from generators import ExportAssets, ExportResult
from models.book_models import BookProject
from services.chapter_writer import author_page_text
from utils.image_utils import script_for
from utils.retry_utils import ExportError, redact
from utils.text_utils import chapter_label, label, parse_blocks, remove_control_chars, split_inline

LOGGER = logging.getLogger("ai_book_studio.docx")

NAVY = RGBColor(0x1E, 0x1B, 0x4B)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
GREY = RGBColor(0x6B, 0x72, 0x80)
_FONT_BY_SCRIPT = {"latin": "Georgia", "tamil": "Nirmala UI", "devanagari": "Nirmala UI", "japanese": "Yu Mincho"}
_PAGE_W, _PAGE_H, _MARGIN = Inches(6), Inches(9), Inches(0.8)


def _set_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    properties = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else None
    if properties is None:
        return
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)


def _field_run(paragraph, kind: str, instruction: str = "") -> None:
    """Append a field character (begin / separate / end) or an instruction run."""
    run = paragraph.add_run()
    if kind == "instr":
        element = OxmlElement("w:instrText")
        element.set(qn("xml:space"), "preserve")
        element.text = instruction
    else:
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
    run._r.append(element)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field_run(paragraph, "begin")
    _field_run(paragraph, "instr", " PAGE ")
    _field_run(paragraph, "separate")
    paragraph.add_run("1")
    _field_run(paragraph, "end")
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = GREY


def _runs(paragraph, text: str) -> None:
    """Add inline Markdown (bold, italic) as runs."""
    for chunk, bold, italic in split_inline(remove_control_chars(text)):
        run = paragraph.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None


def _setup_styles(document, language: str) -> str:
    font = _FONT_BY_SCRIPT[script_for(language)]
    normal = document.styles["Normal"]
    _set_font(normal, font)
    normal.font.size = Pt(11)
    for name, size, colour, before, after in (
        ("Heading 1", 24, NAVY, 0, 18),
        ("Heading 2", 16, NAVY, 14, 6),
        ("Heading 3", 13, INDIGO, 10, 4),
    ):
        style = document.styles[name]
        _set_font(style, font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = colour
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    document.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Page breaks are set on each heading paragraph (see ``_heading``), not on the style, so every viewer starts chapters on a new page.
    document.styles["Heading 1"].paragraph_format.space_before = Pt(60)

    body = document.styles.add_style("BookBody", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.font.size = Pt(11)
    body.paragraph_format.first_line_indent = Inches(0.25)
    body.paragraph_format.space_after = Pt(4)
    body.paragraph_format.line_spacing = 1.25
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if script_for(language) == "latin" else WD_ALIGN_PARAGRAPH.LEFT
    first = document.styles.add_style("BookBodyFirst", WD_STYLE_TYPE.PARAGRAPH)
    first.base_style = body
    first.paragraph_format.first_line_indent = Inches(0)
    return font


def _add_body(document, text: str) -> None:
    fresh = True
    for block in parse_blocks(text):
        if block.kind in {"h1", "h2"}:
            document.add_paragraph(remove_control_chars(block.text), style="Heading 2")
            fresh = True
        elif block.kind == "h3":
            document.add_paragraph(remove_control_chars(block.text), style="Heading 3")
            fresh = True
        elif block.kind == "p":
            _runs(document.add_paragraph(style="BookBodyFirst" if fresh else "BookBody"), block.text)
            fresh = False
        elif block.kind == "bullet":
            _runs(document.add_paragraph(style="List Bullet"), block.text)
        elif block.kind == "number":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.add_run(f"{block.number}. ")
            _runs(paragraph, block.text)
        elif block.kind == "quote":
            paragraph = document.add_paragraph(style="Quote")
            _runs(paragraph, block.text)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(remove_control_chars(block.text))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif block.kind == "break":
            paragraph = document.add_paragraph("*   *   *")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fresh = True


def _add_picture(document, path: Path, max_width: float = 4.4, max_height: float = 3.4) -> None:
    with PILImage.open(path) as picture:
        width, height = picture.size
    inches_w = min(max_width, max_height * width / height)
    document.add_picture(str(path), width=Inches(inches_w))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.paragraphs[-1].paragraph_format.space_after = Pt(12)


class _Flow:
    """Tracks whether the next paragraph is the first one on a fresh page (after a section break)."""

    def __init__(self) -> None:
        self.at_start = True

    def starts_page(self, paragraph) -> None:
        """Make ``paragraph`` start a new page, unless it already begins one (a page break there would leave a blank page)."""
        paragraph.paragraph_format.page_break_before = not self.at_start
        self.at_start = False

    def reset(self) -> None:
        self.at_start = True


def _heading(document, flow: _Flow, text: str):
    paragraph = document.add_paragraph(remove_control_chars(text), style="Heading 1")
    flow.starts_page(paragraph)
    return paragraph


def _section_break(document) -> None:
    """End the current section after the last paragraph, without the empty paragraph python-docx adds for it.

    ``add_section`` appends a new empty paragraph that carries the section properties. That paragraph can spill onto a
    page of its own (a blank page, especially after the full-page cover), so the properties move into the last real paragraph.
    """
    last = document.paragraphs[-1]
    document.add_section(WD_SECTION.NEW_PAGE)
    holder = document.paragraphs[-1]
    properties = holder._p.pPr
    section_properties = properties.find(qn("w:sectPr")) if properties is not None else None
    if section_properties is None or holder._p is last._p:
        return
    last._p.get_or_add_pPr().append(section_properties)
    holder._p.getparent().remove(holder._p)


def _centered(document, text: str, size: int, colour: RGBColor, bold: bool = False, italic: bool = False, before: int = 0, after: int = 6):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(remove_control_chars(text))
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.bold = bold or None
    run.italic = italic or None
    return paragraph


def _add_toc(document, flow: _Flow, entries: list[str], heading: str) -> None:
    """Table of contents field with static entries as fallback (right-click > Update Field in Word adds page numbers)."""
    title = _heading(document, flow, heading)
    # Not a real heading in the TOC: use outline level body text so it is not listed itself.
    title_pr = title._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    title_pr.append(outline)
    if not entries:
        return
    for index, entry in enumerate(entries):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        if index == 0:
            _field_run(paragraph, "begin")
            _field_run(paragraph, "instr", ' TOC \\o "1-1" \\h \\z \\u ')
            _field_run(paragraph, "separate")
        paragraph.add_run(remove_control_chars(entry))
        if index == len(entries) - 1:
            _field_run(paragraph, "end")


def build_docx(project: BookProject, assets: ExportAssets, output_path: Path, include_cover: bool = True) -> ExportResult:
    """Create and validate the DOCX (``include_cover=False`` leaves out the full-page cover). Raises ``ExportError`` on failure."""
    plan = project.plan
    if plan is None:
        raise ExportError("DOCX generation failed: the book has no plan.")
    options = project.book_input
    language = plan.language
    warnings: list[str] = []
    try:
        document = Document()
        _setup_styles(document, language)
        document.core_properties.title = remove_control_chars(plan.title)
        document.core_properties.author = remove_control_chars(plan.author)
        document.core_properties.comments = "Created with AI Book Studio"

        # Cover: its own full-bleed section without margins. The picture is a little smaller than the page and its line
        # is one point high, so the page never overflows into a blank second page.
        flow = _Flow()
        if not include_cover:
            pass  # no cover section: the document starts with the front pages or the body
        elif assets.cover_path and Path(assets.cover_path).exists():
            paragraph = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            with PILImage.open(assets.cover_path) as cover_picture:
                cover_w, cover_h = cover_picture.size
            scale = min(5.9 / cover_w, 8.85 / cover_h)
            run = paragraph.add_run()
            run.font.size = Pt(1)
            run.add_picture(str(assets.cover_path), width=Inches(cover_w * scale), height=Inches(cover_h * scale))
        else:
            paragraph = document.add_paragraph(remove_control_chars(plan.title))
            paragraph.paragraph_format.space_before = Pt(0)
            warnings.append("The DOCX has no cover image.")
        if include_cover:
            _section_break(document)  # the front pages (no page numbers) or the body start in the next section

        has_dedication = options.has("dedication") and bool(project.front_matter.get("dedication", "").strip())
        has_front = options.has("title_page") or options.has("copyright_page") or has_dedication
        if options.has("title_page"):
            flow.starts_page(_centered(document, "", 8, GREY, before=110, after=0))
            _centered(document, plan.title, 30, NAVY, bold=True, before=0, after=10)
            if plan.subtitle:
                _centered(document, plan.subtitle, 15, GREY, italic=True, after=18)
            _centered(document, plan.author, 15, INDIGO, bold=True, before=14)
        if options.has("copyright_page"):
            flow.starts_page(_centered(document, "", 8, GREY, before=260, after=0))
            for line in (
                f"Copyright © {datetime.now().year} {plan.author}",
                "All rights reserved.",
                "This book was created with AI-assisted tools and reviewed by the author.",
            ):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(remove_control_chars(line))
                run.font.size = Pt(9)
                run.font.color.rgb = GREY
        if has_dedication:
            flow.starts_page(_centered(document, "", 8, GREY, before=150, after=0))
            _centered(document, project.front_matter["dedication"].strip(), 12, NAVY, italic=True)
        if has_front:
            _section_break(document)  # the numbered body starts in a new section
            flow.reset()

        def text_section(key: str, text: str) -> None:
            _heading(document, flow, label(key, language))
            _add_body(document, text)

        for key in ("preface", "introduction"):
            if options.has(key) and project.front_matter.get(key, "").strip():
                text_section(key, project.front_matter[key])

        chapters = [c for c in plan.chapters if project.chapters.get(c.chapter_number, "").strip()]
        if options.has("table_of_contents"):
            entries = [f"{chapter_label(c.chapter_number, language)}: {c.title}" for c in chapters]
            _add_toc(document, flow, entries, label("table_of_contents", language))

        for chapter in chapters:
            number = chapter.chapter_number
            _heading(document, flow, f"{chapter_label(number, language)}: {chapter.title}")
            image_path = assets.chapter_images.get(number)
            if image_path and Path(image_path).exists():
                _add_picture(document, Path(image_path))
            _add_body(document, project.chapters[number])

        for key in ("conclusion", "glossary"):
            if options.has(key) and project.front_matter.get(key, "").strip():
                text_section(key, project.front_matter[key])
        if options.has("author_page"):
            text_section("author_page", author_page_text(project))

        # Page set-up for every section: the cover has no margins, the others share the book margins.
        for index, section in enumerate(document.sections):
            section.page_width, section.page_height = _PAGE_W, _PAGE_H
            if index == 0 and include_cover:
                for side in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance"):
                    setattr(section, side, Inches(0))
                continue
            section.left_margin = section.right_margin = _MARGIN
            section.top_margin = Inches(0.9)
            section.bottom_margin = Inches(0.85)
            section.header_distance = Inches(0.4)
            section.footer_distance = Inches(0.45)
        body_section = document.sections[-1]
        body_section.footer.is_linked_to_previous = False
        _add_page_number(body_section.footer.paragraphs[0])  # numbers run through the whole book, cover = page 1 (as in the PDF)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = output_path.with_name(output_path.stem + ".tmp.docx")
        document.save(str(temp_path))
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("DOCX build failed: %s: %s", type(exc).__name__, redact(exc))
        raise ExportError(f"DOCX generation failed: {type(exc).__name__}. See the log for details.") from exc

    try:
        reopened = Document(str(temp_path))
        headings = sum(1 for p in reopened.paragraphs if p.style is not None and p.style.name == "Heading 1")
        if headings < len(chapters):
            raise ExportError("DOCX generation failed: some chapters are missing from the file.")
    except ExportError:
        temp_path.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001
        temp_path.unlink(missing_ok=True)
        raise ExportError("DOCX generation failed: the file could not be read back.") from exc
    temp_path.replace(output_path)
    return ExportResult(path=output_path, pages=0, warnings=warnings)
