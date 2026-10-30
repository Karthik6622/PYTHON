import io
import types

import pytest
from PIL import Image

from src import attachments, image_gen, storage


class FakeUpload:
    def __init__(self, name, data):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


def png_bytes(size=(40, 30), color=(200, 30, 30)):
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, format="PNG")
    return out.getvalue()


def make_docx(paragraphs, table=None):
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pdf(text):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_text_files_are_read():
    assert attachments.read_upload("a.txt", b"hello there")["text"] == "hello there"
    assert attachments.read_upload("a.csv", "name,age\nAsha,30".encode())["kind"] == "document"
    assert "caf\u00e9" in attachments.read_upload("b.md", "caf\u00e9".encode("utf-8"))["text"]


def test_docx_paragraphs_and_tables_are_read():
    item = attachments.read_upload("plan.docx", make_docx(["Launch plan"], [["Task", "Owner"], ["Ship", "Ravi"]]))
    assert "Launch plan" in item["text"] and "Ship | Ravi" in item["text"]


def test_pdf_text_is_read():
    item = attachments.read_upload("doc.pdf", make_pdf("Invoice total 42"))
    assert "Invoice total 42" in item["text"]


def test_bad_files_give_clear_errors():
    with pytest.raises(attachments.AttachmentError, match="empty"):
        attachments.read_upload("a.txt", b"")
    with pytest.raises(attachments.AttachmentError, match="supported"):
        attachments.read_upload("virus.exe", b"MZ")
    with pytest.raises(attachments.AttachmentError, match="at most"):
        attachments.read_upload("big.txt", b"x" * (attachments.MAX_FILE_BYTES + 1))
    with pytest.raises(attachments.AttachmentError, match="PDF"):
        attachments.read_upload("broken.pdf", b"not a pdf")
    with pytest.raises(attachments.AttachmentError, match="Word"):
        attachments.read_upload("broken.docx", b"not a docx")
    with pytest.raises(attachments.AttachmentError, match="no readable text"):
        attachments.read_upload("blank.txt", b"   \n  ")


def test_long_text_is_cut_to_fit():
    text = attachments.read_upload("long.txt", b"a" * 50_000)["text"]
    assert len(text) < attachments.MAX_TEXT_CHARS + 100 and "cut off" in text


def test_file_names_cannot_carry_paths():
    assert attachments.read_upload("..\\..\\evil/notes.txt", b"hi")["name"] == "notes.txt"


def test_image_is_checked_shrunk_and_sent_to_the_vision_model(monkeypatch):
    seen = {}

    class Completions:
        def create(self, **kwargs):
            seen.update(kwargs)
            message = types.SimpleNamespace(content="<think>hmm</think>A red square with the text HELLO.")
            return types.SimpleNamespace(choices=[types.SimpleNamespace(message=message)])

    client = types.SimpleNamespace(chat=types.SimpleNamespace(completions=Completions()))
    monkeypatch.setattr(attachments, "get_model_settings", lambda: object())
    monkeypatch.setattr(attachments, "_openai_client", lambda settings, timeout=None: client)
    item = attachments.read_upload("big.png", png_bytes((3000, 2000)), "what is this?")
    assert item == {"name": "big.png", "kind": "image", "text": "A red square with the text HELLO."}
    assert seen["model"] == "nvidia/nemotron-3-nano-omni-30b-a3b-reasoning"
    content = seen["messages"][0]["content"]
    assert "what is this?" in content[0]["text"]
    assert content[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")


def test_fake_image_is_rejected_before_any_model_call(monkeypatch):
    monkeypatch.setattr(attachments, "get_model_settings", lambda: object())
    monkeypatch.setattr(attachments, "_openai_client", lambda settings, timeout=None: (_ for _ in ()).throw(AssertionError))
    with pytest.raises(attachments.AttachmentError, match="could not be opened"):
        attachments.read_upload("fake.png", b"this is not an image")


def test_vision_timeout_becomes_a_friendly_error(monkeypatch):
    class APITimeoutError(Exception):
        pass

    class Completions:
        def create(self, **kwargs):
            raise APITimeoutError("slow")

    client = types.SimpleNamespace(chat=types.SimpleNamespace(completions=Completions()))
    monkeypatch.setattr(attachments, "get_model_settings", lambda: object())
    monkeypatch.setattr(attachments, "_openai_client", lambda settings, timeout=None: client)
    with pytest.raises(attachments.AttachmentError, match="too long"):
        attachments.read_upload("a.png", png_bytes())


def test_one_bad_upload_does_not_stop_the_others():
    good, problems = attachments.process_uploads(
        [FakeUpload("a.txt", b"one"), FakeUpload("b.exe", b"x"), FakeUpload("c.txt", b"three"), FakeUpload("d.txt", b"four")]
    )
    assert [item["name"] for item in good] == ["a.txt", "c.txt"]
    assert len(problems) == 2 and any("only 3 files" in p for p in problems)


# ---------------------------------------------------------------- image generation

def test_image_requests_are_detected_but_questions_are_not():
    for text in ("generate an image of a red fox", "Create a picture of the sea", "draw a cat", "can you draw me a house",
                 "make a logo for my cafe", "please create a wallpaper with mountains"):
        assert image_gen.wants_image(text), text
    for text in ("how do I create an image in photoshop?", "what is a good picture book", "I took a photo today", "hello"):
        assert not image_gen.wants_image(text), text


def test_image_generation_needs_a_server(monkeypatch):
    monkeypatch.delenv("IMAGE_GEN_BASE_URL", raising=False)
    monkeypatch.setattr(image_gen, "_configured", lambda name, default="": default)
    assert not image_gen.image_generation_configured()
    with pytest.raises(image_gen.ImageGenerationError, match="not set up"):
        image_gen.generate_image("a fox")


def test_image_generation_returns_png(monkeypatch):
    import base64

    settings = {"IMAGE_GEN_BASE_URL": "http://gpu:30000/v1"}
    monkeypatch.setattr(image_gen, "_configured", lambda name, default="": settings.get(name, default))
    seen = {}

    class Images:
        def generate(self, **kwargs):
            seen.update(kwargs)
            return types.SimpleNamespace(data=[types.SimpleNamespace(b64_json=base64.b64encode(png_bytes()).decode())])

    import openai

    monkeypatch.setattr(openai, "OpenAI", lambda **kwargs: types.SimpleNamespace(images=Images()))
    png = image_gen.generate_image("a red fox")
    assert png.startswith(b"\x89PNG")
    assert seen["model"] == "nvidia/Cosmos3-Nano" and seen["prompt"] == "a red fox" and seen["response_format"] == "b64_json"


def test_bad_image_service_reply_is_reported(monkeypatch):
    import base64
    import openai

    monkeypatch.setattr(image_gen, "_configured", lambda name, default="": "http://gpu/v1" if name == "IMAGE_GEN_BASE_URL" else default)
    images = types.SimpleNamespace(generate=lambda **kw: types.SimpleNamespace(data=[types.SimpleNamespace(b64_json=base64.b64encode(b"nope").decode())]))
    monkeypatch.setattr(openai, "OpenAI", lambda **kwargs: types.SimpleNamespace(images=images))
    with pytest.raises(image_gen.ImageGenerationError, match="not a picture"):
        image_gen.generate_image("a fox")


# ---------------------------------------------------------------- storage

def test_messages_keep_attachments_and_generated_images(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path)
    monkeypatch.setattr(storage, "CONVERSATIONS_DIR", tmp_path / "conversations")
    user_id = "a" * 32
    conversation = storage.create_conversation(user_id)
    cid = conversation["conversation_id"]
    storage.append_message(user_id, cid, {
        "role": "user", "content": "read this", "attachments": [{"name": "n.txt", "kind": "document", "text": "x" * 20_000}, "junk"],
    })
    name = storage.save_generated_image(user_id, png_bytes())
    storage.append_message(user_id, cid, {"role": "assistant", "content": "done", "image": name})
    storage.append_message(user_id, cid, {"role": "assistant", "content": "bad", "image": "../../secret.png"})
    saved = storage.load_conversations(user_id)[0]["messages"]
    assert len(saved[0]["attachments"]) == 1 and len(saved[0]["attachments"][0]["text"]) == storage.MAX_ATTACHMENT_TEXT
    assert saved[1]["image"] == name and "image" not in saved[2]
    assert storage.generated_image_bytes(user_id, name).startswith(b"\x89PNG")
    assert storage.generated_image_bytes(user_id, "../../secret.png") is None
    assert storage.generated_image_bytes("b" * 32, name) is None
