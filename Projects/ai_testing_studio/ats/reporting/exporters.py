"""Downloadable report formats: self-contained HTML, Excel workbook, JSON and a ZIP archive with evidence."""

from __future__ import annotations

import base64
import io
import json
import zipfile
from html import escape
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from ats.generation.excel_export import _bytes, _write_table, safe_cell
from ats.security import mask

STATUS_COLORS = {"PASS": "#1b7f3b", "FAIL": "#c62828", "HEALED": "#b26a00", "INCONCLUSIVE": "#6a4bb5", "SKIPPED": "#6b7280", "NOT RUN": "#6b7280"}
MAX_EMBEDDED_IMAGES, MAX_IMAGE_BYTES = 25, 600_000


def report_json(report: dict[str, Any]) -> bytes:
    return json.dumps(report, indent=2, ensure_ascii=False, default=str).encode("utf-8")


# ---------------------------------------------------------------------------------------- Excel
RESULT_COLUMNS = [
    ("Test Case ID", "test_case_id", 16), ("Scenario ID", "scenario_id", 14), ("Module", "module", 18), ("Title", "title", 42), ("Priority", "priority", 10),
    ("Status", "status", 13), ("Duration (s)", "duration", 12), ("Failure Category", "category", 26), ("Confidence", "confidence", 11),
    ("Analysis", "analysis_summary", 50), ("Recommended Action", "recommended_action", 40), ("Error", "message", 60), ("Script", "file", 34),
]  # fmt: skip
DEFECT_COLUMNS = [
    ("Defect ID", "defect_id", 20), ("Test Case ID", "test_case_id", 16), ("Title", "title", 44), ("Severity", "severity", 11), ("Priority", "priority", 11),
    ("Category", "category", 26), ("Description", "description", 56), ("Steps to Reproduce", "steps_to_reproduce", 56), ("Expected Result", "expected_result", 40),
    ("Actual Result", "actual_result", 44), ("Confidence", "confidence", 11), ("Status", "status", 10),
]  # fmt: skip
HEALING_COLUMNS = [
    ("Element", "element_key", 26), ("Test Cases", "test_case_ids", 24), ("Old Locator", "old", 40), ("New Locator", "new", 40), ("Confidence", "confidence", 11),
    ("Status", "status", 12), ("Outcome", "outcome", 22), ("Reason", "reason", 50), ("Pull Request", "pr_url", 40),
]  # fmt: skip
TRACE_COLUMNS = [
    ("Scenario ID", "scenario_id", 14), ("Test Case ID", "test_case_id", 16), ("Module", "module", 18), ("Title", "title", 42), ("Automated", "automated", 10),
    ("Script", "file", 36), ("Steps", "steps", 8), ("Latest Status", "status", 14), ("Notes", "notes", 44),
]  # fmt: skip


def _locator(data: dict[str, Any]) -> str:
    if not data:
        return ""
    name = f" name={data['name']!r}" if data.get("name") else ""
    return f"{data.get('kind', '')}: {data.get('value', '')}{name}"


def report_workbook(report: dict[str, Any]) -> bytes:
    wb, project, execution, totals = Workbook(), report["project"], report["execution"], report["totals"]
    ws = wb.active
    ws.title = "Summary"
    summary = report["summary"]
    rows = [
        ("Project", project["name"]), ("Application URL", project["base_url"]), ("Environment", project["environment"]), ("Run", f"#{execution['run_number']}"),
        ("Execution ID", execution["id"]), ("Framework", execution["framework"]), ("Browser", execution["browser"]), ("Branch", execution["branch"]),
        ("Started", execution["started_at"]), ("Duration (s)", totals.get("duration", 0)), ("Total", totals["total"]), ("Passed", totals["passed"]),
        ("Failed", totals["failed"]), ("Healed", totals["healed"]), ("Inconclusive", totals["inconclusive"]), ("Skipped", totals["skipped"]),
        ("Pass rate %", totals["pass_rate"]), ("", ""), ("Executive summary", summary["summary"]),
        ("Summary source", "AI generated" if summary.get("ai_generated") else "Built from recorded results"),
    ]  # fmt: skip
    rows += [("Highlight", v) for v in summary["highlights"]] + [("Risk", v) for v in summary["risks"]] + [("Recommendation", v) for v in summary["recommendations"]]
    for label, value in rows:
        ws.append([safe_cell(label), safe_cell(value)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
    _write_table(wb.create_sheet("Results"), RESULT_COLUMNS, report["results"])
    _write_table(wb.create_sheet("Traceability"), TRACE_COLUMNS, report["traceability"])
    _write_table(wb.create_sheet("Defects"), DEFECT_COLUMNS, report["defects"])
    healing_rows = [{**h, "old": _locator(h["old_locator"]), "new": _locator(h["new_locator"])} for h in report["healings"]]
    _write_table(wb.create_sheet("Healing"), HEALING_COLUMNS, healing_rows)
    modules = [("Module", "module", 24), ("Total", "total", 9), ("PASS", "PASS", 9), ("FAIL", "FAIL", 9), ("HEALED", "HEALED", 9), ("INCONCLUSIVE", "INCONCLUSIVE", 14), ("SKIPPED", "SKIPPED", 10)]
    _write_table(wb.create_sheet("By Module"), modules, report["by_module"])
    return _bytes(wb)


# ---------------------------------------------------------------------------------------- HTML
CSS = """
body{font-family:Segoe UI,Arial,sans-serif;margin:0;background:#f4f6fa;color:#1f2937}main{max-width:1180px;margin:0 auto;padding:28px}
h1{margin:0 0 4px}h2{margin-top:34px;border-bottom:2px solid #1f3a5f;padding-bottom:4px;color:#1f3a5f}.muted{color:#6b7280}
.cards{display:flex;flex-wrap:wrap;gap:12px;margin:18px 0}.card{background:#fff;border-radius:10px;padding:14px 18px;min-width:120px;box-shadow:0 1px 3px #0002}
.card b{display:block;font-size:26px}table{border-collapse:collapse;width:100%;background:#fff;margin:10px 0;font-size:14px}
th{background:#1f3a5f;color:#fff;text-align:left}td,th{padding:7px 10px;border-bottom:1px solid #e5e7eb;vertical-align:top}
.badge{color:#fff;border-radius:10px;padding:2px 9px;font-size:12px;font-weight:600}img{max-width:100%;border:1px solid #d1d5db;border-radius:6px}
details{background:#fff;border-radius:8px;padding:8px 14px;margin:8px 0;box-shadow:0 1px 2px #0002}summary{cursor:pointer;font-weight:600}pre{white-space:pre-wrap;background:#f3f4f6;padding:8px;border-radius:6px}
"""  # noqa: E501


def _e(value: Any) -> str:
    return escape(mask(str(value if value is not None else "")))


def _badge(status: str) -> str:
    return f'<span class="badge" style="background:{STATUS_COLORS.get(status, "#374151")}">{_e(status)}</span>'


def _list(items: list[str]) -> str:
    return "<ul>" + "".join(f"<li>{_e(i)}</li>" for i in items) + "</ul>" if items else '<p class="muted">None.</p>'


def _table(headers: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{_e(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>" for row in rows)
    return f"<table><tr>{head}</tr>{body}</table>"


def report_html(report: dict[str, Any], evidence: dict[str, bytes] | None = None) -> bytes:
    """A single file with no external resources. All values are escaped; screenshots of failures are embedded."""
    evidence = evidence or {}
    project, execution, totals, summary = report["project"], report["execution"], report["totals"], report["summary"]
    cards = "".join(
        f'<div class="card"><span class="muted">{label}</span><b>{value}</b></div>'
        for label, value in (("Total", totals["total"]), ("Passed", totals["passed"]), ("Failed", totals["failed"]), ("Healed", totals["healed"]),
                             ("Inconclusive", totals["inconclusive"]), ("Skipped", totals["skipped"]), ("Pass rate", f"{totals['pass_rate']}%"))
    )  # fmt: skip
    parts = [
        f"<h1>{_e(project['name'])} - Test Execution Report</h1>",
        f'<p class="muted">Run #{_e(execution["run_number"])} | {_e(execution["framework"])} on {_e(execution["browser"])} | {_e(project["environment"])} | '
        f'{_e(project["base_url"])} | {_e(execution["started_at"])} | generated {_e(report["generated_at"])}</p>',
        f'<div class="cards">{cards}</div>',
        "<h2>Executive summary</h2>",
        f"<p>{_e(summary['summary'])}</p>",
        f'<p class="muted">{"Written by the AI model from the recorded results." if summary.get("ai_generated") else "Built directly from the recorded results."}</p>',
        "<h3>Highlights</h3>" + _list(summary["highlights"]), "<h3>Risks</h3>" + _list(summary["risks"]), "<h3>Recommendations</h3>" + _list(summary["recommendations"]),
        "<p class=\"muted\">HEALED means a test passed only after a locator fix proposed by the system; it is counted separately from PASS. "
        "INCONCLUSIVE means a failure did not reproduce on re-check.</p>",
    ]  # fmt: skip
    if report["by_module"]:
        parts.append("<h2>Results by module</h2>" + _table(["Module", "Total", "PASS", "FAIL", "HEALED", "INCONCLUSIVE", "SKIPPED"], [[_e(m["module"])] + [_e(m[k]) for k in ("total", "PASS", "FAIL", "HEALED", "INCONCLUSIVE", "SKIPPED")] for m in report["by_module"]]))
    if report["by_category"]:
        parts.append("<h2>Failure categories</h2>" + _table(["Category", "Count"], [[_e(k), _e(v)] for k, v in sorted(report["by_category"].items())]))
    parts.append("<h2>Test results</h2>" + _table(["Test case", "Title", "Module", "Status", "Duration", "Category"], [[_e(r["test_case_id"]), _e(r["title"]), _e(r["module"]), _badge(r["status"]), f"{r['duration']}s", _e(r["category"])] for r in report["results"]]))

    failures = [r for r in report["results"] if r["status"] in {"FAIL", "INCONCLUSIVE", "HEALED"}]
    if failures:
        parts.append("<h2>Failure and healing details</h2>")
        embedded = 0
        for r in failures:
            body = [f"<p>{_badge(r['status'])} {_e(r['category'])} {('(confidence ' + _e(r['confidence']) + ')') if r['confidence'] != '' else ''}</p>"]
            if r["analysis_summary"]:
                body.append(f"<p><b>Analysis:</b> {_e(r['analysis_summary'])}</p>")
            if r["recommended_action"]:
                body.append(f"<p><b>Recommended action:</b> {_e(r['recommended_action'])}</p>")
            if r["message"]:
                body.append(f"<pre>{_e(r['message'])}</pre>")
            shot = evidence.get(r["screenshot"]) if r["screenshot"] else None
            if shot and embedded < MAX_EMBEDDED_IMAGES and len(shot) <= MAX_IMAGE_BYTES:
                embedded += 1
                body.append(f'<img alt="Screenshot for {_e(r["test_case_id"])}" src="data:image/png;base64,{base64.b64encode(shot).decode()}">')
            parts.append(f"<details><summary>{_e(r['test_case_id'])} - {_e(r['title'])}</summary>{''.join(body)}</details>")
    if report["healings"]:
        parts.append("<h2>Self-healing</h2>" + _table(["Element", "Old locator", "New locator", "Confidence", "Status", "Outcome"], [[_e(h["element_key"]), _e(_locator(h["old_locator"])), _e(_locator(h["new_locator"])), _e(f"{h['confidence']:.2f}"), _e(h["status"]), _e(h["outcome"])] for h in report["healings"]]))
    if report["defects"]:
        parts.append("<h2>Draft defects</h2>")
        for d in report["defects"]:
            steps = "".join(f"<li>{_e(s)}</li>" for s in d.get("steps_to_reproduce", []))
            parts.append(
                f"<details><summary>{_e(d.get('defect_id'))} - {_e(d.get('title'))} ({_e(d.get('severity'))})</summary><p>{_e(d.get('description'))}</p>"
                f"<p><b>Steps</b></p><ol>{steps}</ol><p><b>Expected:</b> {_e(d.get('expected_result'))}</p><p><b>Actual:</b> {_e(d.get('actual_result'))}</p></details>"
            )  # fmt: skip
    parts.append("<h2>Traceability</h2>" + _table(["Scenario", "Test case", "Title", "Script", "Status"], [[_e(t["scenario_id"]), _e(t["test_case_id"]), _e(t["title"]), _e(t["file"]), _badge(t["status"])] for t in report["traceability"]]))
    if report["history"]:
        parts.append("<h2>Recent runs</h2>" + _table(["Run", "Status", "Passed", "Failed", "Healed", "Pass rate"], [[_e(h["run_number"]), _e(h["status"]), _e(h["totals"].get("passed", "")), _e(h["totals"].get("failed", "")), _e(h["totals"].get("healed", "")), _e(h["totals"].get("pass_rate", ""))] for h in report["history"]]))
    document = f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'"><title>{_e(project["name"])} report</title><style>{CSS}</style></head><body><main>{"".join(parts)}</main></body></html>'
    return document.encode("utf-8")


# ---------------------------------------------------------------------------------------- ZIP
def report_zip(report: dict[str, Any], evidence: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("report.html", report_html(report, evidence))
        z.writestr("report.xlsx", report_workbook(report))
        z.writestr("report.json", report_json(report))
        for name, data in evidence.items():
            z.writestr(name, data)
    return buffer.getvalue()
