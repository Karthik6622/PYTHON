"""Report assembly and exports."""

from __future__ import annotations

import io
import json
import re
import zipfile

import openpyxl

from ats.ai.agents import ReportAgent
from ats.errors import ModelError
from ats.security import register_secret
from ats.reporting import ReportService, deterministic_summary, report_html, report_json, report_workbook, report_zip
from tests.fakes import FakeClient
from tests.test_analysis import TC, build_service, classification, healing_reply, make_failed_execution

SUMMARY = {"summary": "Run completed with one healed test.", "highlights": ["Login recovered"], "risks": ["Button was renamed"], "recommendations": ["Update the page object"]}


def summary_client() -> FakeClient:
    return FakeClient({"(defect mode)": {"title": "t", "description": "d", "severity": "High", "priority": "High", "actual_result": "a"}, "Report Agent": SUMMARY})


class Down(FakeClient):
    def complete(self, *_a, **_k):
        raise ModelError("model unavailable")


def healed_run(repo):
    service, executor, _client, _ = build_service(repo, classification(), healing_reply(), "PASS")
    project, ctx = make_failed_execution(repo, executor, error="Timeout 30000ms exceeded waiting for get_by_role('button', name='Sign in')")
    service.policies[project.id] = "auto_pr"
    service.post_run(ctx)
    return project, ctx


def test_report_contains_counts_healing_and_ai_summary(repo):
    project, ctx = healed_run(repo)
    report = ReportService(repo, ReportAgent(summary_client(), repo)).build(project, ctx.execution_id)
    assert report["totals"]["healed"] == 1 and report["totals"]["passed"] == 0
    assert report["results"][0]["status"] == "HEALED" and report["results"][0]["category"] == "LOCATOR DRIFT"
    assert report["healings"][0]["element_key"] == "login.sign_in_button" and report["healings"][0]["status"] in {"verified", "applied"}
    assert report["summary"]["ai_generated"] is True and report["summary"]["summary"] == SUMMARY["summary"]
    assert report["history"] == [] or isinstance(report["history"], list)


def test_summary_falls_back_when_model_is_down(repo):
    project, ctx = healed_run(repo)
    report = ReportService(repo, ReportAgent(Down(), repo)).build(project, ctx.execution_id)
    assert report["summary"]["ai_generated"] is False and "1 healed" in report["summary"]["summary"]
    assert deterministic_summary(report)["summary"] == report["summary"]["summary"]


def test_exports_are_valid_and_neutralise_hostile_content(repo):
    project, ctx = healed_run(repo)
    service = ReportService(repo, ReportAgent(Down(), repo))
    register_secret("hunter2xyz-secret")
    repo.update_result(ctx.execution_id, TC, message="=cmd|' /C calc'!A0 <script>alert(1)</script> password hunter2xyz-secret", status="FAIL")
    report = service.build(project, ctx.execution_id, with_ai=False)
    evidence = service.evidence_files(project.id, repo.get_execution(ctx.execution_id))
    assert any(name.endswith("steps.json") for name in evidence) and "reports/evidence/TC_LOG_001/failure.json" in evidence

    html = report_html(report, evidence).decode()
    assert "<script>alert" not in html and "&lt;script&gt;" in html and "http-equiv=\"Content-Security-Policy\"" in html
    assert not re.search(r"(?:src|href)=\"https?://", html)

    workbook = openpyxl.load_workbook(io.BytesIO(report_workbook(report)))
    assert {"Summary", "Results", "Traceability", "Defects", "Healing", "By Module"} <= set(workbook.sheetnames)
    cells = [str(c.value) for row in workbook["Results"].iter_rows() for c in row if c.value]
    assert not any(v.startswith("=") for v in cells)

    assert json.loads(report_json(report))["project"]["name"] == "Shop"
    with zipfile.ZipFile(io.BytesIO(report_zip(report, evidence))) as z:
        names = set(z.namelist())
        assert {"report.html", "report.xlsx", "report.json", "run.log"} - names <= {"run.log"}
        assert not any(n.endswith("dom.html") for n in names)
        assert all("hunter2xyz-secret" not in z.read(n).decode("utf-8", "ignore") for n in ("report.json", "report.html"))
